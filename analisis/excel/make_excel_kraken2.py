import csv, glob, math, os, re, shutil
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.workbook.properties import CalcProperties

BASE = r"C:\Users\zamir\Downloads\Libro_Proyecto_Grado"
E = BASE + r"\datos_energia"
OUTDIR = BASE + r"\excel graficas"
OUT = OUTDIR + r"\graficas_kraken2.xlsx"

CL_DIR = E + r"\cluster_rpi5\kraken2"
ZA_RUN = E + r"\x86_pc_zamir\kraken2\energia_kraken2_x86_lote_run_20260929_123305.csv"
ZA_PRE = E + r"\x86_pc_zamir\kraken2\pre\energia_kraken2_x86_lote_pre_20260929_123305.csv"
ZA_MARCAS = E + r"\x86_pc_zamir\kraken2\marcas_por_muestra_20260929_123305.csv"
CR_RUN = E + r"\x86_pc_cristian\kraken2\energia_kraken2_Nativo_all_20260929_112010_20260929_112613.csv"
CR_PRE = E + r"\x86_pc_cristian\kraken2\pre\energia_baseline_pre_kraken2_20260929_112110.csv"
SAMPLES = ["c001", "c002", "c003", "e001", "e002"]
CL, ZA, CR = "Clúster", "Laptop Zamir", "Laptop Cristian"

# viejo -> excel graficas\viejos
old_dir = os.path.join(OUTDIR, "viejos")
os.makedirs(old_dir, exist_ok=True)
VIEJO = os.path.join(old_dir, "graficas_kraken2_pre30s.xlsx")  # versión con las corridas de pre 30 s
if os.path.exists(OUT) and not os.path.exists(VIEJO):
    shutil.move(OUT, VIEJO)

F = "Calibri"
f_n = Font(name=F, size=11)
f_b = Font(name=F, size=11, bold=True)
f_t = Font(name=F, size=13, bold=True)
f_in = Font(name=F, size=11, color="0000FF")
f_g = Font(name=F, size=10, color="666666")
head = PatternFill("solid", fgColor="E7E6E6")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
TS = "yyyy-mm-dd hh:mm:ss.000"
C_CR, C_ZA, C_CL = "2A78D6", "EB6834", "1BAF7A"
C_NODE = {"c001": "2A78D6", "c002": "EB6834", "c003": "1BAF7A", "e001": "EDA100", "e002": "8E5CC4"}


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def ts(s):
    return datetime.strptime(s, "%Y-%m-%d %H:%M:%S.%f")


def header(ws, names, widths, row=1, col=1):
    for i, (n, w) in enumerate(zip(names, widths), col):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill, c.border = f_b, head, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32


# ---------------- leer datos ----------------
cl_run, cl_pre, cl_node = {}, {}, {}
for s in SAMPLES:
    fr = glob.glob(os.path.join(CL_DIR, f"energia_Kraken2_multinodo_{s}_*.csv"))[0]
    fp = glob.glob(os.path.join(CL_DIR, "pre", f"energia_baseline_pre_kraken2_{s}_*.csv"))[0]
    cl_node[s] = re.search(r"node_(rbp5-\d)", fr).group(1)
    cl_run[s] = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(fr)]
    cl_pre[s] = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(fp)]
za_run = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(ZA_RUN)]
za_pre = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(ZA_PRE)]
za_marcas = [(r["muestra"], ts(r["inicio"]), ts(r["fin"])) for r in read(ZA_MARCAS)]
cr_run = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(CR_RUN)]
cr_pre = [(ts(r["timestamp"]), float(r["power_w"])) for r in read(CR_PRE)]

wb = Workbook()

# ---------------- LEEME ----------------
ws = wb.active
ws.title = "LEEME"
ws.column_dimensions["A"].width = 110
lines = [
    ("Kraken 2 - potencia durante la corrida", f_t),
    ("", f_n),
    ("Datos", f_b),
    ("Clúster (5 Raspberry Pi 5): las 5 muestras corren en paralelo, una por nodo. Carpeta cluster_rpi5\\kraken2.", f_n),
    ("Laptop Zamir (x86, WSL2): lote de 5 muestras seguidas. Las marcas de inicio y fin de cada muestra están en Datos_Zamir.", f_n),
    ("Laptop Cristian (x86, Ubuntu nativo): lote de 5 muestras seguidas, sin separar por muestra.", f_n),
    ("", f_n),
    ("Baseline y normalización", f_b),
    ("Antes de cada corrida hay 60 s de estabilización sin medir y 5 min de reposo medido (pre).", f_n),
    ("El baseline es el promedio de ese pre, uno por nodo en el clúster.", f_n),
    ("Potencia normalizada = potencia medida - baseline, cada plataforma con el suyo.", f_n),
    ("Energía = integral trapezoidal con los tiempos reales de cada lectura, en Wh (hoja Baselines).", f_n),
    ("", f_n),
    ("Hojas", f_b),
    ("Datos_Cluster, Datos_Zamir, Datos_Cristian: lecturas de la corrida. Datos_Pre: lecturas del reposo.", f_n),
    ("Azul = dato del CSV, negro = fórmula. t (s) = tiempo desde el inicio de la corrida; el pre queda en tiempo negativo.", f_n),
    ("G1 a G3: la corrida y, encima, la curva real del reposo (pre), las dos desde 0 s para compararlas.", f_n),
    ("El eje llega hasta el final de la corrida; el reposo completo (5 min) está en Datos_Pre.", f_n),
    ("t reposo (s) en Datos_Pre = tiempo desde la primera lectura del reposo.", f_n),
    ("G1: cada nodo del clúster con sus lecturas reales; los reposos de los 5 nodos van en gris.", f_n),
    ("G4: potencia normalizada de las tres plataformas. Clúster = suma de los 5 nodos interpolada cada 1 s,", f_n),
    ("porque cada nodo lee en momentos distintos (tabla a la derecha de la gráfica).", f_n),
    ("", f_n),
    ("Nota: en el clúster el medidor leyó de forma irregular durante Kraken (14 a 15 lecturas en unos 50 s);", f_n),
    ("las lecturas reales de cada nodo están en Datos_Cluster.", f_n),
]
for i, (t, fnt) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = fnt

# ---------------- Baselines (se crea antes para fijar filas) ----------------
wp = wb.create_sheet("Datos_Pre")
wbl = wb.create_sheet("Baselines")
wcl = wb.create_sheet("Datos_Cluster")
wza = wb.create_sheet("Datos_Zamir")
wcr = wb.create_sheet("Datos_Cristian")

header(wbl, ["Plataforma", "Muestra", "Nodo", "Lecturas pre", "Baseline (W)", "Inicio corrida"],
       [16, 10, 10, 11, 13, 23])
base_rows = [(CL, s, cl_node[s]) for s in SAMPLES] + [(ZA, "lote", ""), (CR, "lote", "")]
BL = {}
for i, (eq, s, nodo) in enumerate(base_rows, 2):
    wbl.cell(row=i, column=1, value=eq)
    wbl.cell(row=i, column=2, value=s)
    wbl.cell(row=i, column=3, value=nodo)
    wbl.cell(row=i, column=4, value=f"=COUNTIFS(Datos_Pre!$A:$A,A{i},Datos_Pre!$B:$B,B{i})")
    wbl.cell(row=i, column=5, value=f"=AVERAGEIFS(Datos_Pre!$E:$E,Datos_Pre!$A:$A,A{i},Datos_Pre!$B:$B,B{i})")
    if eq == CL:
        start = f"=_xlfn.MINIFS(Datos_Cluster!$C:$C,Datos_Cluster!$A:$A,B{i})"
    else:
        start = "=MIN(Datos_Zamir!$B:$B)" if eq == ZA else "=MIN(Datos_Cristian!$B:$B)"
    wbl.cell(row=i, column=6, value=start)
    wbl.cell(row=i, column=5).number_format = "0.000"
    wbl.cell(row=i, column=6).number_format = TS
    for c in range(1, 7):
        wbl.cell(row=i, column=c).font = f_n
        wbl.cell(row=i, column=c).border = box
    BL[(eq, s)] = f"Baselines!$E${i}"

# ---------------- Datos_Pre ----------------
header(wp, ["Plataforma", "Muestra", "Hora", "t (s)", "Potencia (W)", "Baseline (W)", "Potencia normalizada (W)",
             "t reposo (s)"],
       [16, 10, 23, 9, 13, 13, 15, 10])
wp.freeze_panes = "A2"
pre_rows = [(CL, s, t, p) for s in SAMPLES for t, p in cl_pre[s]]
pre_rows += [(ZA, "lote", t, p) for t, p in za_pre]
pre_rows += [(CR, "lote", t, p) for t, p in cr_pre]
rpre = {}
for i, (eq, s, t, p) in enumerate(pre_rows, 2):
    for c, v in ((1, eq), (2, s), (3, t), (5, p)):
        wp.cell(row=i, column=c, value=v).font = f_in
    wp.cell(row=i, column=4, value=f"=ROUND((C{i}-SUMIFS(Baselines!$F:$F,Baselines!$A:$A,A{i},Baselines!$B:$B,B{i}))*86400,1)")
    wp.cell(row=i, column=6, value=f"=SUMIFS(Baselines!$E:$E,Baselines!$A:$A,A{i},Baselines!$B:$B,B{i})")
    wp.cell(row=i, column=7, value=f"=E{i}-F{i}")
    wp.cell(row=i, column=8, value=f"=ROUND((C{i}-_xlfn.MINIFS($C:$C,$A:$A,A{i},$B:$B,B{i}))*86400,1)")
    for c in (4, 6, 7, 8):
        wp.cell(row=i, column=c).font = f_n
    wp.cell(row=i, column=3).number_format = TS
    for c in (5, 6, 7):
        wp.cell(row=i, column=c).number_format = "0.000"
    lo, hi = rpre.get((eq, s), (i, i))
    rpre[(eq, s)] = (min(lo, i), max(hi, i))

# ---------------- Datos_Cluster ----------------
header(wcl, ["Muestra", "Nodo", "Hora", "t (s)", "t lote (s)", "Potencia (W)", "Baseline (W)",
             "Potencia normalizada (W)", "Energía tramo (Wh)", "Energía neta tramo (Wh)"],
       [10, 10, 23, 9, 11, 13, 13, 15, 13, 14])
wcl.freeze_panes = "A2"
rcl = {}
i = 2
for s in SAMPLES:
    first = i
    for t, p in cl_run[s]:
        wcl.cell(row=i, column=1, value=s).font = f_in
        wcl.cell(row=i, column=2, value=cl_node[s]).font = f_in
        wcl.cell(row=i, column=3, value=t).font = f_in
        wcl.cell(row=i, column=6, value=p).font = f_in
        wcl.cell(row=i, column=4, value=f"=ROUND((C{i}-_xlfn.MINIFS($C:$C,$A:$A,A{i}))*86400,1)")
        wcl.cell(row=i, column=5, value=f"=ROUND((C{i}-MIN($C:$C))*86400,1)")
        wcl.cell(row=i, column=7, value=f"={BL[(CL, s)]}")
        wcl.cell(row=i, column=8, value=f"=F{i}-G{i}")
        if i > first:
            wcl.cell(row=i, column=9, value=f"=(F{i}+F{i-1})/2*(C{i}-C{i-1})*24")
            wcl.cell(row=i, column=10, value=f"=(H{i}+H{i-1})/2*(C{i}-C{i-1})*24")
        for c in (4, 5, 7, 8, 9, 10):
            wcl.cell(row=i, column=c).font = f_n
        wcl.cell(row=i, column=3).number_format = TS
        for c in (6, 7, 8):
            wcl.cell(row=i, column=c).number_format = "0.000"
        for c in (9, 10):
            wcl.cell(row=i, column=c).number_format = "0.00000"
        i += 1
    rcl[s] = (first, i - 1)


# ---------------- Datos_Zamir / Datos_Cristian ----------------
def lote_sheet(w, rows, eq, muestra_formula):
    header(w, ["Muestra", "Hora", "t (s)", "Potencia (W)", "Baseline (W)", "Potencia normalizada (W)",
               "Energía tramo (Wh)", "Energía neta tramo (Wh)"], [16, 23, 9, 13, 13, 15, 13, 14])
    w.freeze_panes = "A2"
    for i, (t, p) in enumerate(rows, 2):
        if muestra_formula:
            w.cell(row=i, column=1, value=muestra_formula(i)).font = f_n
        else:
            w.cell(row=i, column=1, value="lote").font = f_in
        w.cell(row=i, column=2, value=t).font = f_in
        w.cell(row=i, column=4, value=p).font = f_in
        w.cell(row=i, column=3, value=f"=ROUND((B{i}-MIN($B:$B))*86400,1)")
        w.cell(row=i, column=5, value=f"={BL[(eq, 'lote')]}")
        w.cell(row=i, column=6, value=f"=D{i}-E{i}")
        if i > 2:
            w.cell(row=i, column=7, value=f"=(D{i}+D{i-1})/2*(B{i}-B{i-1})*24")
            w.cell(row=i, column=8, value=f"=(F{i}+F{i-1})/2*(B{i}-B{i-1})*24")
        for c in (3, 5, 6, 7, 8):
            w.cell(row=i, column=c).font = f_n
        w.cell(row=i, column=2).number_format = TS
        for c in (4, 5, 6):
            w.cell(row=i, column=c).number_format = "0.000"
        for c in (7, 8):
            w.cell(row=i, column=c).number_format = "0.00000"
    return 2, len(rows) + 1


# marcas de Zamir en J:L
header(wza, ["Muestra", "Inicio", "Fin"], [10, 23, 23], row=1, col=10)
for k, (s, a, b) in enumerate(za_marcas, 2):
    wza.cell(row=k, column=10, value=s).font = f_in
    wza.cell(row=k, column=11, value=a).font = f_in
    wza.cell(row=k, column=12, value=b).font = f_in
    wza.cell(row=k, column=11).number_format = TS
    wza.cell(row=k, column=12).number_format = TS
mk = f"$K$2:$K${len(za_marcas) + 1}"
ml = f"$L$2:$L${len(za_marcas) + 1}"
mj = f"$J$2:$J${len(za_marcas) + 1}"


def za_muestra(i):
    return (f'=IF(B{i}<MIN({mk}),"fuera de muestra",IF(B{i}>INDEX({ml},MATCH(B{i},{mk},1)),'
            f'"fuera de muestra",INDEX({mj},MATCH(B{i},{mk},1))))')


rza = lote_sheet(wza, za_run, ZA, za_muestra)
rcr = lote_sheet(wcr, cr_run, CR, None)

# ---------------- energía (bajo Baselines) ----------------
r0 = len(base_rows) + 3
wbl.cell(row=r0, column=1, value="Energía de la corrida").font = f_b
header(wbl, ["Plataforma", "Lecturas", "Duración (s)", "Energía bruta (Wh)", "Energía neta (Wh)"],
       [16, 10, 10, 11, 13], row=r0 + 1)
wbl.column_dimensions["A"].width = 16
en = [
    (CL, "=COUNT(Datos_Cluster!F:F)", "=MAX(Datos_Cluster!E:E)", "=SUM(Datos_Cluster!I:I)", "=SUM(Datos_Cluster!J:J)"),
    (ZA, "=COUNT(Datos_Zamir!D:D)", "=MAX(Datos_Zamir!C:C)", "=SUM(Datos_Zamir!G:G)", "=SUM(Datos_Zamir!H:H)"),
    (CR, "=COUNT(Datos_Cristian!D:D)", "=MAX(Datos_Cristian!C:C)", "=SUM(Datos_Cristian!G:G)", "=SUM(Datos_Cristian!H:H)"),
]
for k, row in enumerate(en, r0 + 2):
    for c, v in enumerate(row, 1):
        cell = wbl.cell(row=k, column=c, value=v)
        cell.font, cell.border = f_n, box
    wbl.cell(row=k, column=3).number_format = "0.0"
    wbl.cell(row=k, column=4).number_format = "0.0000"
    wbl.cell(row=k, column=5).number_format = "0.0000"
wbl.cell(row=r0 + 5, column=1, value="Clúster = suma de los 5 nodos; duración = nodo más largo.").font = f_g


# ---------------- gráficas ----------------
def series(ws, lo, hi, xcol, ycol, title, color, dash=False, marker=True, width=19050, msize=5):
    s = Series(Reference(ws, min_col=ycol, min_row=lo, max_row=hi),
               Reference(ws, min_col=xcol, min_row=lo, max_row=hi), title=title)
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width
    if dash:
        s.graphicalProperties.line.dashStyle = "sysDash"
    if marker:
        s.marker.symbol = "circle"
        s.marker.size = msize
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        s.marker.graphicalProperties.line.solidFill = color
    else:
        s.marker.symbol = "none"
    s.smooth = False
    return s


def chart(title, xt, yt, xmin, xmax, xunit, ymin, ymax, yunit):
    c = ScatterChart()
    c.title, c.style = title, 13
    c.x_axis.title, c.y_axis.title = xt, yt
    c.x_axis.delete = c.y_axis.delete = False
    c.x_axis.majorGridlines = None
    c.title.overlay = False
    c.legend.position = "r"
    c.legend.overlay = False
    for ax, lo, hi, u in ((c.x_axis, xmin, xmax, xunit), (c.y_axis, ymin, ymax, yunit)):
        ax.number_format = "0"
        ax.numFmt.sourceLinked = False
        ax.scaling.min, ax.scaling.max, ax.majorUnit = lo, hi, u
        ax.crosses = "min"
    c.width, c.height = 26, 13
    return c


def hide(c, idxs):
    c.legend.legendEntry = [LegendEntry(idx=k, delete=True) for k in idxs]


def note(ws, text):
    ws["A1"] = text
    ws["A1"].font = f_g


XMIN = -35

C_REP = "8C8C8C"  # reposo en gris para distinguirlo de la corrida
TC = 18  # columna R: tablas auxiliares a la derecha de las gráficas
L = lambda c: chr(64 + c)


def interp_formula(sheet, tcol, ncol, lo, hi, xcell):
    """interpolación lineal de la columna ncol en el tiempo xcell, sobre las filas lo..hi."""
    T = f"{sheet}!${tcol}${lo}:${tcol}${hi}"
    N = f"{sheet}!${ncol}${lo}:${ncol}${hi}"
    k = f"MATCH({xcell},{T},1)"
    return (f"=IFERROR(INDEX({N},{k})+(INDEX({N},{k}+1)-INDEX({N},{k}))*({xcell}-INDEX({T},{k}))"
            f"/(INDEX({T},{k}+1)-INDEX({T},{k})),INDEX({N},{k}))")


# G1 clúster: los 5 nodos con sus lecturas reales y, en gris, el reposo de cada nodo; todo desde 0 s
g1 = wb.create_sheet("G1_Cluster_vs_baseline")
note(g1, "Corrida de Kraken 2 en cada nodo (lecturas reales) y, en gris, el reposo previo de cada nodo, desde 0 s.")
c = chart("Kraken 2 en el clúster: potencia por nodo vs. reposo", "Tiempo (s)", "Potencia (W)", 0, 55, 5, 1, 7, 1)
for s in SAMPLES:
    lo, hi = rcl[s]
    c.series.append(series(wcl, lo, hi, 4, 6, f"{s} ({cl_node[s]})", C_NODE[s]))
G1_HIDE = []
for k, s in enumerate(SAMPLES):
    lo, hi = rpre[(CL, s)]
    if k:
        G1_HIDE.append(len(c.series) + 1)  # 1-based, para ocultarlos en la leyenda desde Excel
    c.series.append(series(wp, lo, hi, 8, 5, "Reposo (cada nodo)" if k == 0 else f"Reposo {s}", C_REP,
                           width=12700, msize=3))
g1.add_chart(c, "A3")
print("G1_HIDE=" + ",".join(map(str, G1_HIDE)))


def lote_chart(ws_g, ws_d, rng, eq, color, title, label, xmax, ymax, yunit):
    c = chart(title, "Tiempo (s)", "Potencia (W)", 0, xmax, 5, 0, ymax, yunit)
    lo, hi = rng
    c.series.append(series(ws_d, lo, hi, 3, 4, label, color))
    lo, hi = rpre[(eq, "lote")]
    c.series.append(series(wp, lo, hi, 8, 5, "Reposo (pre)", C_REP, msize=4))
    ws_g.add_chart(c, "A3")


g2 = wb.create_sheet("G2_Zamir_vs_baseline")
note(g2, "Lote de 5 muestras de Kraken 2 en la laptop de Zamir (WSL2) y su reposo previo, las dos desde 0 s.")
lote_chart(g2, wza, rza, ZA, C_ZA, "Kraken 2 en laptop Zamir (WSL2): potencia vs. reposo",
           "Corrida (lote 5 muestras)", 55, 50, 10)

g3 = wb.create_sheet("G3_Cristian_vs_baseline")
note(g3, "Lote de 5 muestras de Kraken 2 en la laptop de Cristian (nativo) y su reposo previo, las dos desde 0 s.")
lote_chart(g3, wcr, rcr, CR, C_CR, "Kraken 2 en laptop Cristian (nativo): potencia vs. reposo",
           "Corrida (lote 5 muestras)", 25, 50, 10)

# G4: clúster = suma de los 5 nodos normalizados, interpolada a una rejilla de 1 s (fórmulas)
g4 = wb.create_sheet("G4_Comparacion")
note(g4, "Potencia normalizada (medida - baseline). Clúster = suma de los 5 nodos interpolada cada 1 s (tabla de la derecha). "
         "Tiempo negativo = reposo.")
t0 = {s: cl_run[s][0][0] for s in SAMPLES}
pre_t = {s: [round((t - t0[s]).total_seconds(), 1) for t, _ in cl_pre[s]] for s in SAMPLES}
run_t = {s: [round((t - t0[s]).total_seconds(), 1) for t, _ in cl_run[s]] for s in SAMPLES}
pre_grid = list(range(math.ceil(max(v[0] for v in pre_t.values())), math.floor(min(v[-1] for v in pre_t.values())) + 1))
run_grid = list(range(0, math.floor(min(v[-1] for v in run_t.values())) + 1))
header(g4, ["t (s)"] + [f"{s} norm. (W)" for s in SAMPLES] + ["Clúster suma (W)"], [7] + [10] * 5 + [11], row=3, col=TC)


r = 4
g4_pre = (r, r + len(pre_grid) - 1)
g4_run = (r + len(pre_grid), r + len(pre_grid) + len(run_grid) - 1)
for x in pre_grid + run_grid:
    xc = f"${chr(64 + TC)}{r}"
    g4.cell(row=r, column=TC, value=x).font = f_n
    for k, s in enumerate(SAMPLES):
        if x < 0:
            lo, hi = rpre[(CL, s)]
            f = interp_formula("Datos_Pre", "D", "G", lo, hi, xc)
        else:
            lo, hi = rcl[s]
            f = interp_formula("Datos_Cluster", "D", "H", lo, hi, xc)
        g4.cell(row=r, column=TC + 1 + k, value=f).font = f_n
        g4.cell(row=r, column=TC + 1 + k).number_format = "0.000"
    g4.cell(row=r, column=TC + 6, value=f"=SUM({chr(65 + TC)}{r}:{chr(69 + TC)}{r})").font = f_n
    g4.cell(row=r, column=TC + 6).number_format = "0.000"
    r += 1
zr = r + 1
lo0, hi0, xc0, yc0 = zr + 1, zr + 2, TC, TC + 1
g4.cell(row=zr, column=TC, value="Línea 0 W").font = f_b
for k, xv in enumerate((XMIN, 60)):
    g4.cell(row=zr + 1 + k, column=TC, value=xv).font = f_n
    g4.cell(row=zr + 1 + k, column=TC + 1, value=0).font = f_in

c = chart("Kraken 2: potencia normalizada por plataforma", "Tiempo desde el inicio de la corrida (s)",
          "Potencia normalizada (W)", XMIN, 60, 5, -10, 50, 10)
c.series.append(series(g4, *g4_run, TC, TC + 6, "Clúster (suma 5 nodos)", C_CL))
c.series.append(series(wza, *rza, 3, 6, "Laptop Zamir (WSL2)", C_ZA))
c.series.append(series(wcr, *rcr, 3, 6, "Laptop Cristian (nativo)", C_CR))
c.series.append(series(g4, *g4_pre, TC, TC + 6, "Clúster pre", C_CL))
c.series.append(series(wp, *rpre[(ZA, "lote")], 4, 7, "Zamir pre", C_ZA))
c.series.append(series(wp, *rpre[(CR, "lote")], 4, 7, "Cristian pre", C_CR))
c.series.append(series(g4, lo0, hi0, xc0, yc0, "Baseline (0 W)", "404040", dash=True, marker=False, width=9525))
hide(c, [3, 4, 5])
g4.add_chart(c, "A3")

for w in wb.worksheets:
    for row in w.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name != F:
                cell.font = Font(name=F, size=cell.font.size, bold=cell.font.bold, color=cell.font.color)

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("OK", OUT)
