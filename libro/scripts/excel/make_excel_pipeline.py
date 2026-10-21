"""Excel de energía por pipeline (LotuS3, QIIME 2) con la misma estructura que graficas_kraken2.xlsx.
Uso: python make_excel_pipeline.py lotus3|qiime2"""
import csv, glob, math, os, re, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.workbook.properties import CalcProperties

BASE = r"C:\Users\zamir\Downloads\Libro_Proyecto_Grado"
E = BASE + r"\datos_energia"
OUTDIR = BASE + r"\excel graficas"

CFG = {
    "lotus3": dict(
        nombre="LotuS3", archivo="graficas_lotus3.xlsx",
        cl_run=r"\cluster_rpi5\lotus3\energia_LotuS3_multinodo_{s}_*.csv",
        cl_pre=r"\cluster_rpi5\lotus3\pre\energia_baseline_pre_lotus3_{s}_*.csv",
        samples=["c001", "c002", "e001"],
        za_run=r"\x86_pc_zamir\lotus3\energia_lotus3_x86_lote_run_20260929_002718.csv",
        za_pre=r"\x86_pc_zamir\lotus3\pre\energia_lotus3_x86_lote_pre_20260929_002718.csv",
        za_marcas=r"\x86_pc_zamir\lotus3\marcas_por_muestra_20260929_002718.csv",
        cr_run=r"\x86_pc_cristian\lotus3\energia_LotuS3_Nativo_Batch3_SILVA_20260927_135552.csv",
        cr_pre=r"\x86_pc_cristian\lotus3\pre\energia_baseline_pre_LotuS3_20260927_135522.csv",
        unidad="s", marcadores=True,
        # (xmin, xmax, xunit, ymin, ymax, yunit)
        ejes={"G1": (-35, 35, 5, 1, 7, 1), "G2": (-40, 90, 10, 0, 50, 10),
              "G3": (-35, 65, 5, 0, 40, 10), "G4": (-40, 90, 10, -10, 30, 10)},
        nota_cl="En el clúster el medidor leyó casi cada segundo (intervalo máximo de 1.3 s).",
    ),
    "qiime2": dict(
        nombre="QIIME 2", archivo="graficas_qiime2.xlsx",
        cl_run=r"\cluster_rpi5\qiime2\energia_Qiime2_multinodo_{s}_*.csv",
        cl_pre=r"\cluster_rpi5\qiime2\pre\energia_baseline_pre_qiime2_{s}_*.csv",
        samples=["c001", "c002", "c003", "e001", "e002"],
        za_run=r"\x86_pc_zamir\qiime2\energia_qiime2_x86_lote_run_20260929_003500.csv",
        za_pre=r"\x86_pc_zamir\qiime2\pre\energia_qiime2_x86_lote_pre_20260929_003500.csv",
        za_marcas=r"\x86_pc_zamir\qiime2\marcas_por_muestra_20260929_003500.csv",
        cr_run=r"\x86_pc_cristian\qiime2\energia_Qiime2_Nativo_Batch5_SILVA_20260918_125756.csv",
        cr_pre=r"\x86_pc_cristian\qiime2\pre\energia_baseline_pre_Qiime2_20260918_125726.csv",
        unidad="min", marcadores=False,
        ejes={"G1": (-5, 30, 5, 1, 6, 1), "G2": (-10, 90, 10, 0, 50, 10),
              "G3": (-10, 70, 10, 0, 50, 10), "G4": (-10, 90, 10, -20, 50, 10)},
        nota_cl="En el clúster el medidor tuvo huecos de hasta unos 8 s entre lecturas durante la corrida.",
    ),
}

P = CFG[sys.argv[1]]
OUT = os.path.join(OUTDIR, P["archivo"])
SAMPLES = P["samples"]
CL, ZA, CR = "Clúster", "Laptop Zamir", "Laptop Cristian"
MIN = P["unidad"] == "min"
FACT, DEC = (1440, 3) if MIN else (86400, 1)   # días -> unidad de tiempo, decimales de t
TU = "t (min)" if MIN else "t (s)"
TL = "t lote (min)" if MIN else "t lote (s)"
XT = f"Tiempo desde el inicio de la corrida ({P['unidad']})"
NN = P["nombre"]

F = "Calibri"
f_n = Font(name=F, size=11)
f_b = Font(name=F, size=11, bold=True)
f_t = Font(name=F, size=13, bold=True)
f_in = Font(name=F, size=11, color="0000FF")
f_g = Font(name=F, size=10, color="666666")
head = PatternFill("solid", fgColor="E7E6E6")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
TS = "yyyy-mm-dd hh:mm:ss.000"
TFMT = "0.000" if MIN else "0.0"
C_CR, C_ZA, C_CL = "2A78D6", "EB6834", "1BAF7A"
C_NODE = {"c001": "2A78D6", "c002": "EB6834", "c003": "1BAF7A", "e001": "EDA100", "e002": "8E5CC4"}


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def ts(s):
    return datetime.strptime(s[:23], "%Y-%m-%d %H:%M:%S.%f")


def pts(path):
    return [(ts(r["timestamp"]), float(r["power_w"])) for r in read(path)]


def header(ws, names, widths, row=1, col=1):
    for i, (n, w) in enumerate(zip(names, widths), col):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill, c.border = f_b, head, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32


# ---------------- leer datos ----------------
cl_run, cl_pre, cl_node = {}, {}, {}
for s in SAMPLES:
    fr = glob.glob(E + P["cl_run"].format(s=s))[0]
    cl_node[s] = re.search(r"node_(rbp5-\d)", fr).group(1)
    cl_run[s] = pts(fr)
    cl_pre[s] = pts(glob.glob(E + P["cl_pre"].format(s=s))[0])
za_run, za_pre = pts(E + P["za_run"]), pts(E + P["za_pre"])
za_marcas = [(r["muestra"], ts(r["inicio"]), ts(r["fin"])) for r in read(E + P["za_marcas"])]
cr_run, cr_pre = pts(E + P["cr_run"]), pts(E + P["cr_pre"])
n_cl = len(SAMPLES)

wb = Workbook()

# ---------------- LEEME ----------------
ws = wb.active
ws.title = "LEEME"
ws.column_dimensions["A"].width = 110
lines = [
    (f"{NN} - potencia durante la corrida", f_t),
    ("", f_n),
    ("Datos", f_b),
    (f"Clúster ({n_cl} Raspberry Pi 5): las {n_cl} muestras corren en paralelo, una por nodo. Carpeta cluster_rpi5\\{sys.argv[1]}.", f_n),
    (f"Laptop Zamir (x86, WSL2): lote de {n_cl} muestras seguidas. Las marcas de inicio y fin de cada muestra están en Datos_Zamir.", f_n),
    (f"Laptop Cristian (x86, Ubuntu nativo): lote de {n_cl} muestras seguidas, sin separar por muestra.", f_n),
    ("", f_n),
    ("Baseline y normalización", f_b),
    ("Antes de cada corrida hay 30 s de reposo (pre). El baseline es el promedio de ese pre, uno por nodo en el clúster.", f_n),
    ("Potencia normalizada = potencia medida - baseline, cada plataforma con el suyo.", f_n),
    ("Energía = integral trapezoidal con los tiempos reales de cada lectura, en Wh (hoja Baselines).", f_n),
    ("", f_n),
    ("Hojas", f_b),
    ("Datos_Cluster, Datos_Zamir, Datos_Cristian: lecturas de la corrida. Datos_Pre: lecturas del reposo.", f_n),
    (f"Azul = dato del CSV, negro = fórmula. {TU} = tiempo desde el inicio de la corrida; el pre queda en tiempo negativo.", f_n),
    ("G1 a G3: curva del pre y de la corrida; la línea punteada es el promedio del pre (baseline).", f_n),
    (f"G4: potencia normalizada de las tres plataformas. Clúster = suma de los {n_cl} nodos interpolada cada 1 s,", f_n),
    ("porque cada nodo lee en momentos distintos (tabla a la derecha de la gráfica). Un nodo que ya terminó suma 0.", f_n),
    ("", f_n),
    ("Nota: " + P["nota_cl"], f_n),
]
if MIN:
    lines.append(("Las gráficas van en minutos; el pre (unos 0.5 min) se ve como un tramo corto antes de 0.", f_n))
for i, (t, fnt) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = fnt

wp = wb.create_sheet("Datos_Pre")
wbl = wb.create_sheet("Baselines")
wcl = wb.create_sheet("Datos_Cluster")
wza = wb.create_sheet("Datos_Zamir")
wcr = wb.create_sheet("Datos_Cristian")

# ---------------- Baselines ----------------
header(wbl, ["Plataforma", "Muestra", "Nodo", "Lecturas pre", "Baseline (W)", "Inicio corrida"],
       [16, 10, 10, 11, 13, 23])
base_rows = [(CL, s, cl_node[s]) for s in SAMPLES] + [(ZA, "lote", ""), (CR, "lote", "")]
BL, T0 = {}, {}
for i, (eq, s, nodo) in enumerate(base_rows, 2):
    wbl.cell(row=i, column=1, value=eq)
    wbl.cell(row=i, column=2, value=s)
    wbl.cell(row=i, column=3, value=nodo)
    wbl.cell(row=i, column=4, value=f"=COUNTIFS(Datos_Pre!$A:$A,A{i},Datos_Pre!$B:$B,B{i})")
    wbl.cell(row=i, column=5, value=f"=AVERAGEIFS(Datos_Pre!$E:$E,Datos_Pre!$A:$A,A{i},Datos_Pre!$B:$B,B{i})")
    if eq == CL:
        start = f"=_xlfn.MINIFS(Datos_Cluster!$C:$C,Datos_Cluster!$A:$A,B{i})"
    else:
        start = "=MIN(Datos_Zamir!$B:$B)" if eq == ZA else "=MIN(Datos_Cristian!$B:$B)"
    wbl.cell(row=i, column=6, value=start)
    wbl.cell(row=i, column=5).number_format = "0.000"
    wbl.cell(row=i, column=6).number_format = TS
    for c in range(1, 7):
        wbl.cell(row=i, column=c).font = f_n
        wbl.cell(row=i, column=c).border = box
    BL[(eq, s)] = f"Baselines!$E${i}"
    T0[(eq, s)] = f"Baselines!$F${i}"
cl_first_row = 2
T0_CL_ALL = f"MIN(Baselines!$F${cl_first_row}:$F${cl_first_row + n_cl - 1})"

# ---------------- Datos_Pre ----------------
header(wp, ["Plataforma", "Muestra", "Hora", TU, "Potencia (W)", "Baseline (W)", "Potencia normalizada (W)"],
       [16, 10, 23, 9, 13, 13, 15])
wp.freeze_panes = "A2"
pre_rows = [(CL, s, t, p) for s in SAMPLES for t, p in cl_pre[s]]
pre_rows += [(ZA, "lote", t, p) for t, p in za_pre]
pre_rows += [(CR, "lote", t, p) for t, p in cr_pre]
rpre = {}
for i, (eq, s, t, p) in enumerate(pre_rows, 2):
    for c, v in ((1, eq), (2, s), (3, t), (5, p)):
        wp.cell(row=i, column=c, value=v).font = f_in
    wp.cell(row=i, column=4, value=f"=ROUND((C{i}-{T0[(eq, s)]})*{FACT},{DEC})")
    wp.cell(row=i, column=6, value=f"={BL[(eq, s)]}")
    wp.cell(row=i, column=7, value=f"=E{i}-F{i}")
    for c in (4, 6, 7):
        wp.cell(row=i, column=c).font = f_n
    wp.cell(row=i, column=3).number_format = TS
    wp.cell(row=i, column=4).number_format = TFMT
    for c in (5, 6, 7):
        wp.cell(row=i, column=c).number_format = "0.000"
    lo, hi = rpre.get((eq, s), (i, i))
    rpre[(eq, s)] = (min(lo, i), max(hi, i))

# ---------------- Datos_Cluster ----------------
header(wcl, ["Muestra", "Nodo", "Hora", TU, TL, "Potencia (W)", "Baseline (W)",
             "Potencia normalizada (W)", "Energía tramo (Wh)", "Energía neta tramo (Wh)"],
       [10, 10, 23, 9, 11, 13, 13, 15, 13, 14])
wcl.freeze_panes = "A2"
rcl = {}
i = 2
for s in SAMPLES:
    first = i
    for t, p in cl_run[s]:
        wcl.cell(row=i, column=1, value=s).font = f_in
        wcl.cell(row=i, column=2, value=cl_node[s]).font = f_in
        wcl.cell(row=i, column=3, value=t).font = f_in
        wcl.cell(row=i, column=6, value=p).font = f_in
        wcl.cell(row=i, column=4, value=f"=ROUND((C{i}-{T0[(CL, s)]})*{FACT},{DEC})")
        wcl.cell(row=i, column=5, value=f"=ROUND((C{i}-{T0_CL_ALL})*{FACT},{DEC})")
        wcl.cell(row=i, column=7, value=f"={BL[(CL, s)]}")
        wcl.cell(row=i, column=8, value=f"=F{i}-G{i}")
        if i > first:
            wcl.cell(row=i, column=9, value=f"=(F{i}+F{i-1})/2*(C{i}-C{i-1})*24")
            wcl.cell(row=i, column=10, value=f"=(H{i}+H{i-1})/2*(C{i}-C{i-1})*24")
        for c in (4, 5, 7, 8, 9, 10):
            wcl.cell(row=i, column=c).font = f_n
        wcl.cell(row=i, column=3).number_format = TS
        wcl.cell(row=i, column=4).number_format = TFMT
        wcl.cell(row=i, column=5).number_format = TFMT
        for c in (6, 7, 8):
            wcl.cell(row=i, column=c).number_format = "0.000"
        for c in (9, 10):
            wcl.cell(row=i, column=c).number_format = "0.00000"
        i += 1
    rcl[s] = (first, i - 1)


# ---------------- Datos_Zamir / Datos_Cristian ----------------
def lote_sheet(w, rows, eq, muestra_formula):
    header(w, ["Muestra", "Hora", TU, "Potencia (W)", "Baseline (W)", "Potencia normalizada (W)",
               "Energía tramo (Wh)", "Energía neta tramo (Wh)"], [16, 23, 9, 13, 13, 15, 13, 14])
    w.freeze_panes = "A2"
    for i, (t, p) in enumerate(rows, 2):
        if muestra_formula:
            w.cell(row=i, column=1, value=muestra_formula(i)).font = f_n
        else:
            w.cell(row=i, column=1, value="lote").font = f_in
        w.cell(row=i, column=2, value=t).font = f_in
        w.cell(row=i, column=4, value=p).font = f_in
        w.cell(row=i, column=3, value=f"=ROUND((B{i}-{T0[(eq, 'lote')]})*{FACT},{DEC})")
        w.cell(row=i, column=5, value=f"={BL[(eq, 'lote')]}")
        w.cell(row=i, column=6, value=f"=D{i}-E{i}")
        if i > 2:
            w.cell(row=i, column=7, value=f"=(D{i}+D{i-1})/2*(B{i}-B{i-1})*24")
            w.cell(row=i, column=8, value=f"=(F{i}+F{i-1})/2*(B{i}-B{i-1})*24")
        for c in (3, 5, 6, 7, 8):
            w.cell(row=i, column=c).font = f_n
        w.cell(row=i, column=2).number_format = TS
        w.cell(row=i, column=3).number_format = TFMT
        for c in (4, 5, 6):
            w.cell(row=i, column=c).number_format = "0.000"
        for c in (7, 8):
            w.cell(row=i, column=c).number_format = "0.00000"
    return 2, len(rows) + 1


header(wza, ["Muestra", "Inicio", "Fin"], [10, 23, 23], row=1, col=10)
for k, (s, a, b) in enumerate(za_marcas, 2):
    wza.cell(row=k, column=10, value=s).font = f_in
    wza.cell(row=k, column=11, value=a).font = f_in
    wza.cell(row=k, column=12, value=b).font = f_in
    wza.cell(row=k, column=11).number_format = TS
    wza.cell(row=k, column=12).number_format = TS
mk = f"$K$2:$K${len(za_marcas) + 1}"
ml = f"$L$2:$L${len(za_marcas) + 1}"
mj = f"$J$2:$J${len(za_marcas) + 1}"


def za_muestra(i):
    return (f'=IF(B{i}<MIN({mk}),"fuera de muestra",IF(B{i}>INDEX({ml},MATCH(B{i},{mk},1)),'
            f'"fuera de muestra",INDEX({mj},MATCH(B{i},{mk},1))))')


rza = lote_sheet(wza, za_run, ZA, za_muestra)
rcr = lote_sheet(wcr, cr_run, CR, None)

# ---------------- energía (bajo Baselines) ----------------
r0 = len(base_rows) + 3
wbl.cell(row=r0, column=1, value="Energía de la corrida").font = f_b
DUR = "Duración (min)" if MIN else "Duración (s)"
header(wbl, ["Plataforma", "Lecturas", DUR, "Energía bruta (Wh)", "Energía neta (Wh)"],
       [16, 10, 10, 11, 13], row=r0 + 1)
en = [
    (CL, "=COUNT(Datos_Cluster!F:F)", "=MAX(Datos_Cluster!E:E)", "=SUM(Datos_Cluster!I:I)", "=SUM(Datos_Cluster!J:J)"),
    (ZA, "=COUNT(Datos_Zamir!D:D)", "=MAX(Datos_Zamir!C:C)", "=SUM(Datos_Zamir!G:G)", "=SUM(Datos_Zamir!H:H)"),
    (CR, "=COUNT(Datos_Cristian!D:D)", "=MAX(Datos_Cristian!C:C)", "=SUM(Datos_Cristian!G:G)", "=SUM(Datos_Cristian!H:H)"),
]
for k, row in enumerate(en, r0 + 2):
    for c, v in enumerate(row, 1):
        cell = wbl.cell(row=k, column=c, value=v)
        cell.font, cell.border = f_n, box
    wbl.cell(row=k, column=3).number_format = "0.0"
    wbl.cell(row=k, column=4).number_format = "0.0000"
    wbl.cell(row=k, column=5).number_format = "0.0000"
wbl.cell(row=r0 + 5, column=1,
         value=f"Clúster = suma de los {n_cl} nodos; duración = desde el primer nodo que arranca hasta el último que termina.").font = f_g


# ---------------- gráficas ----------------
def series(ws, lo, hi, xcol, ycol, title, color, dash=False, marker=None, width=19050, msize=5):
    marker = P["marcadores"] if marker is None else marker
    s = Series(Reference(ws, min_col=ycol, min_row=lo, max_row=hi),
               Reference(ws, min_col=xcol, min_row=lo, max_row=hi), title=title)
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width
    if dash:
        s.graphicalProperties.line.dashStyle = "sysDash"
    if marker:
        s.marker.symbol = "circle"
        s.marker.size = msize
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        s.marker.graphicalProperties.line.solidFill = color
    else:
        s.marker.symbol = "none"
    s.smooth = False
    return s


def chart(title, yt, key):
    xmin, xmax, xunit, ymin, ymax, yunit = P["ejes"][key]
    c = ScatterChart()
    c.title, c.style = title, 13
    c.x_axis.title, c.y_axis.title = XT, yt
    c.x_axis.delete = c.y_axis.delete = False
    c.x_axis.majorGridlines = None
    c.title.overlay = False
    c.legend.position = "r"
    c.legend.overlay = False
    for ax, lo, hi, u in ((c.x_axis, xmin, xmax, xunit), (c.y_axis, ymin, ymax, yunit)):
        ax.number_format = "0"
        ax.numFmt.sourceLinked = False
        ax.scaling.min, ax.scaling.max, ax.majorUnit = lo, hi, u
        ax.crosses = "min"
    c.width, c.height = 26, 13
    return c


def helper_line(ws, row, col, label, x0, x1, yref):
    ws.cell(row=row, column=col, value=label).font = f_b
    ws.cell(row=row + 1, column=col, value=x0).font = f_n
    ws.cell(row=row + 2, column=col, value=x1).font = f_n
    ws.cell(row=row + 1, column=col + 1, value=f"={yref}").font = f_n
    ws.cell(row=row + 2, column=col + 1, value=f"={yref}").font = f_n
    ws.cell(row=row + 1, column=col + 1).number_format = "0.000"
    ws.cell(row=row + 2, column=col + 1).number_format = "0.000"
    return row + 1, row + 2, col, col + 1


def note(ws, text):
    ws["A1"] = text
    ws["A1"].font = f_g


def to_u(sec):
    return sec / 60 if MIN else sec


hide_idx = {}

# G1 clúster
g1 = wb.create_sheet("G1_Cluster_vs_baseline")
note(g1, f"Reposo (tiempo negativo) y corrida de {NN} en cada nodo. Línea punteada = promedio del pre de cada nodo.")
xmin1, xend1 = P["ejes"]["G1"][0], P["ejes"]["G1"][1]
c = chart(f"{NN} en el clúster: potencia por nodo vs. baseline", "Potencia (W)", "G1")
hidden = []
for s in SAMPLES:
    lo, hi = rcl[s]
    c.series.append(series(wcl, lo, hi, 4, 6, f"{s} ({cl_node[s]})", C_NODE[s]))
for s in SAMPLES:
    lo, hi = rpre[(CL, s)]
    hidden.append(len(c.series))
    c.series.append(series(wp, lo, hi, 4, 5, f"{s} pre", C_NODE[s]))
for k, s in enumerate(SAMPLES):
    lo, hi, xc, yc = helper_line(g1, 3 + 4 * k, 18, f"Baseline {s}", xmin1, xend1, BL[(CL, s)])
    if k:
        hidden.append(len(c.series))
    c.series.append(series(g1, lo, hi, xc, yc, "Promedio pre (cada nodo)" if k == 0 else f"Prom. {s}",
                           "404040", dash=True, marker=False, width=9525))
hide_idx["G1_Cluster_vs_baseline"] = hidden
g1.add_chart(c, "A3")


def lote_chart(name, ws_d, rng, eq, color, title, label, key):
    ws_g = wb.create_sheet(name)
    note(ws_g, f"Reposo (tiempo negativo) y lote de {n_cl} muestras de {NN} en la {label.split(' (')[0].lower()}. "
               "Línea punteada = promedio del pre.")
    c = chart(title, "Potencia (W)", key)
    lo, hi = rng
    c.series.append(series(ws_d, lo, hi, 3, 4, label, color))
    lo, hi = rpre[(eq, "lote")]
    c.series.append(series(wp, lo, hi, 4, 5, "Reposo (pre)", color))
    lo, hi, xc, yc = helper_line(ws_g, 3, 18, "Baseline", P["ejes"][key][0], P["ejes"][key][1], BL[(eq, "lote")])
    c.series.append(series(ws_g, lo, hi, xc, yc, "Promedio pre (baseline)", "404040", dash=True, marker=False, width=9525))
    hide_idx[name] = [1]
    ws_g.add_chart(c, "A3")


lote_chart("G2_Zamir_vs_baseline", wza, rza, ZA, C_ZA, f"{NN} en laptop Zamir (WSL2): potencia vs. baseline",
           f"Laptop Zamir (lote {n_cl} muestras)", "G2")
lote_chart("G3_Cristian_vs_baseline", wcr, rcr, CR, C_CR, f"{NN} en laptop Cristian (nativo): potencia vs. baseline",
           f"Laptop Cristian (lote {n_cl} muestras)", "G3")

# G4: clúster = suma de nodos normalizados, interpolada cada 1 s (fórmulas); un nodo que ya terminó suma 0
g4 = wb.create_sheet("G4_Comparacion")
note(g4, f"Potencia normalizada (medida - baseline). Clúster = suma de los {n_cl} nodos interpolada cada 1 s "
         "(tabla de la derecha). Tiempo negativo = reposo.")
t0 = {s: cl_run[s][0][0] for s in SAMPLES}
pre_t = {s: [(t - t0[s]).total_seconds() for t, _ in cl_pre[s]] for s in SAMPLES}
run_end = {s: (cl_run[s][-1][0] - t0[s]).total_seconds() for s in SAMPLES}
pre_grid = list(range(math.ceil(max(v[0] for v in pre_t.values())), math.floor(min(v[-1] for v in pre_t.values())) + 1))
run_grid = list(range(0, math.floor(max(run_end.values())) + 1))
TC = 18
header(g4, [TU] + [f"{s} norm. (W)" for s in SAMPLES] + ["Clúster suma (W)"], [8] + [10] * n_cl + [11], row=3, col=TC)
L = lambda c: chr(64 + c)


def interp_formula(sheet, tcol, ncol, lo, hi, xcell, end_zero):
    T = f"{sheet}!${tcol}${lo}:${tcol}${hi}"
    N = f"{sheet}!${ncol}${lo}:${ncol}${hi}"
    k = f"MATCH({xcell},{T},1)"
    f = (f"IFERROR(INDEX({N},{k})+(INDEX({N},{k}+1)-INDEX({N},{k}))*({xcell}-INDEX({T},{k}))"
         f"/(INDEX({T},{k}+1)-INDEX({T},{k})),INDEX({N},{k}))")
    return f"=IF({xcell}>MAX({T}),0,{f})" if end_zero else "=" + f


r = 4
g4_pre = (r, r + len(pre_grid) - 1)
g4_run = (r + len(pre_grid), r + len(pre_grid) + len(run_grid) - 1)
for x in pre_grid + run_grid:
    xc = f"${L(TC)}{r}"
    g4.cell(row=r, column=TC, value=round(to_u(x), 4)).font = f_n
    g4.cell(row=r, column=TC).number_format = TFMT
    for k, s in enumerate(SAMPLES):
        if x < 0:
            lo, hi = rpre[(CL, s)]
            f = interp_formula("Datos_Pre", "D", "G", lo, hi, xc, False)
        else:
            lo, hi = rcl[s]
            f = interp_formula("Datos_Cluster", "D", "H", lo, hi, xc, True)
        g4.cell(row=r, column=TC + 1 + k, value=f).font = f_n
        g4.cell(row=r, column=TC + 1 + k).number_format = "0.000"
    g4.cell(row=r, column=TC + 1 + n_cl, value=f"=SUM({L(TC + 1)}{r}:{L(TC + n_cl)}{r})").font = f_n
    g4.cell(row=r, column=TC + 1 + n_cl).number_format = "0.000"
    r += 1
zr = r + 1
g4.cell(row=zr, column=TC, value="Línea 0 W").font = f_b
for k, xv in enumerate((P["ejes"]["G4"][0], P["ejes"]["G4"][1])):
    g4.cell(row=zr + 1 + k, column=TC, value=xv).font = f_n
    g4.cell(row=zr + 1 + k, column=TC + 1, value=0).font = f_in

c = chart(f"{NN}: potencia normalizada por plataforma", "Potencia normalizada (W)", "G4")
SUMC = TC + 1 + n_cl
run_series = [series(g4, *g4_run, TC, SUMC, f"Clúster (suma {n_cl} nodos)", C_CL),
              series(wza, *rza, 3, 6, "Laptop Zamir (WSL2)", C_ZA),
              series(wcr, *rcr, 3, 6, "Laptop Cristian (nativo)", C_CR)]
if MIN:  # con miles de puntos, el clúster se dibuja al final para que quede encima
    run_series = run_series[1:] + run_series[:1]
c.series.extend(run_series)
c.series.append(series(g4, *g4_pre, TC, SUMC, "Clúster pre", C_CL))
c.series.append(series(wp, *rpre[(ZA, "lote")], 4, 7, "Zamir pre", C_ZA))
c.series.append(series(wp, *rpre[(CR, "lote")], 4, 7, "Cristian pre", C_CR))
c.series.append(series(g4, zr + 1, zr + 2, TC, TC + 1, "Baseline (0 W)", "404040", dash=True, marker=False, width=9525))
hide_idx["G4_Comparacion"] = [3, 4, 5]
g4.add_chart(c, "A3")

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
# índices 1-based de leyenda a ocultar, para el paso en Excel
print("HIDE " + ";".join(f"{k}=" + ",".join(str(v + 1) for v in vs) for k, vs in hide_idx.items()))
print("OK", OUT)
